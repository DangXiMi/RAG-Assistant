"""Live end-to-end test against a running API + worker.

Uploads one file per supported format through POST /api/v1/ingest, waits for
the arq worker to finish, then queries the retrieval pipeline and checks that
answers are grounded in the uploaded documents.

Usage:
    python scripts/e2e_live_test.py [base_url] [--fresh]
"""
from __future__ import annotations

import io
import os
import sys
import tempfile
import time
from pathlib import Path

import requests

BASE_URL = sys.argv[1] if len(sys.argv) > 1 else "http://127.0.0.1:8010"
# Pass --fresh to clear Redis dedup keys first, so a re-run indexes from scratch
# instead of reporting every chunk as already indexed.
FRESH = "--fresh" in sys.argv
POLL_TIMEOUT = 300

PASS, FAIL, WARN = "PASS", "FAIL", "WARN"
results: list[tuple[str, str, str]] = []


def record(name: str, status: str, detail: str = "") -> None:
    results.append((name, status, detail))
    icon = {PASS: "[ok]", FAIL: "[XX]", WARN: "[!!]"}[status]
    print(f"{icon} {name}" + (f" -- {detail}" if detail else ""), flush=True)


def check(name: str, condition: bool, detail: str = "") -> None:
    record(name, PASS if condition else FAIL, detail)


# --- Fixture generation -----------------------------------------------------
# Every document answers the same domain question so retrieval is comparable.

REFUND_FACT = "Refunds are processed within 30 days of delivery."


def build_files(workdir: Path) -> list[Path]:
    made: list[Path] = []

    p = workdir / "refund_policy.txt"
    p.write_text(
        "Refund Policy\n\n"
        f"{REFUND_FACT}\n"
        "Shipping costs are non-refundable for change-of-mind returns.\n",
        encoding="utf-8",
    )
    made.append(p)

    p = workdir / "refund_policy.html"
    p.write_text(
        "<html><body><h1>Refund Policy</h1>"
        f"<p>{REFUND_FACT}</p>"
        "<table><tr><th>Item</th><th>Window</th></tr>"
        "<tr><td>Sale item</td><td>Not refundable</td></tr></table>"
        "</body></html>",
        encoding="utf-8",
    )
    made.append(p)

    p = workdir / "refund_policy.csv"
    p.write_text(
        "Item type,Refund window\n"
        "Standard,30 days from delivery\n"
        "Sale item,Not refundable\n",
        encoding="utf-8",
    )
    made.append(p)

    from docx import Document

    p = workdir / "refund_policy.docx"
    document = Document()
    document.add_heading("Refund Policy", level=1)
    document.add_paragraph(REFUND_FACT)
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Item type"
    table.cell(0, 1).text = "Refund window"
    table.cell(1, 0).text = "Sale item"
    table.cell(1, 1).text = "Not refundable"
    document.save(p)
    made.append(p)

    from openpyxl import Workbook

    p = workdir / "refund_matrix.xlsx"
    workbook = Workbook()
    workbook.remove(workbook.active)
    policy = workbook.create_sheet("Policy")
    policy.append(["Item type", "Refund window"])
    policy.append(["Standard", "30 days"])
    policy.append(["Sale item", "Not refundable"])
    shipping = workbook.create_sheet("Shipping")
    shipping.append(["Region", "Lead time"])
    shipping.append(["Europe", "5 business days"])
    workbook.save(p)
    made.append(p)

    import pymupdf

    p = workdir / "refund_policy.pdf"
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 100), REFUND_FACT, fontsize=12)
    page.insert_text((72, 130), "Page one contains the general policy.", fontsize=12)
    second = document.new_page()
    second.insert_text((72, 100), "Express shipping is non-refundable.", fontsize=12)
    document.save(str(p))
    document.close()
    made.append(p)

    from PIL import Image, ImageDraw

    p = workdir / "scanned_refund.png"
    image = Image.new("RGB", (1100, 240), "white")
    ImageDraw.Draw(image).text(
        (20, 100), "REFUNDS WITHIN 30 DAYS OF DELIVERY", fill="black", font_size=52
    )
    image.save(p)
    made.append(p)

    # A PDF with no text layer at all: must be read by OCR.
    p = workdir / "scanned_policy.pdf"
    document = pymupdf.open()
    scan_page = document.new_page(width=1100, height=240)
    scan_page.insert_image(pymupdf.Rect(0, 0, 1100, 240), filename=str(workdir / "scanned_refund.png"))
    document.save(str(p))
    document.close()
    made.append(p)

    return made


# --- API helpers ------------------------------------------------------------

def ingest(path: Path) -> dict:
    with open(path, "rb") as handle:
        response = requests.post(
            f"{BASE_URL}/api/v1/ingest",
            files={"file": (path.name, handle, "application/octet-stream")},
            timeout=120,
        )
    response.raise_for_status()
    return response.json()


def wait_for_job(job_id: str) -> dict:
    deadline = time.time() + POLL_TIMEOUT
    while time.time() < deadline:
        response = requests.get(f"{BASE_URL}/api/v1/job/{job_id}", timeout=30)
        if response.status_code == 200:
            data = response.json()
            if data.get("status") in ("done", "failed"):
                return data
        time.sleep(3)
    return {"status": "timeout", "result": None}


def query(question: str, mode: str = "Hybrid") -> dict:
    response = requests.post(
        f"{BASE_URL}/api/v1/query",
        json={"question": question, "user_id": "e2e", "mode": mode},
        timeout=300,
    )
    response.raise_for_status()
    return response.json()


# --- Main -------------------------------------------------------------------

def main() -> int:
    print("=" * 74)
    print(f"LIVE END-TO-END TEST against {BASE_URL}")
    print("=" * 74)

    try:
        health = requests.get(f"{BASE_URL}/api/v1/health", timeout=10).json()
        check("API reachable", health.get("status") == "ok", str(health))
    except Exception as exc:
        record("API reachable", FAIL, f"{type(exc).__name__}: {exc}")
        return 1

    formats = requests.get(f"{BASE_URL}/api/v1/formats", timeout=10).json()
    check("OCR reported available", formats["ocr"]["available"] is True,
          str(formats["ocr"].get("reason")))

    # Clear the content-hash dedup cache so this run indexes from scratch.
    # Without this, a second run of this script finds every chunk already
    # indexed and legitimately ingests nothing.
    if FRESH:
        try:
            import redis as redis_lib

            client = redis_lib.Redis(
                host=os.getenv("REDIS_HOST", "localhost"),
                port=int(os.getenv("REDIS_PORT", "6379")),
                decode_responses=True,
            )
            keys = list(client.scan_iter("chunk:hash:*", count=1000))
            if keys:
                client.delete(*keys)
            record("dedup cache cleared", PASS, f"{len(keys)} key(s) removed")
        except Exception as exc:
            record("dedup cache cleared", WARN, f"{type(exc).__name__}: {exc}")

    print()
    print("--- Ingestion through the API + arq worker ---")

    with tempfile.TemporaryDirectory() as tmp:
        files = build_files(Path(tmp))
        ingested: list[tuple[str, dict]] = []

        for path in files:
            try:
                payload = ingest(path)
            except Exception as exc:
                record(f"{path.name}: upload", FAIL, f"{type(exc).__name__}: {exc}")
                continue

            job = wait_for_job(payload["job_id"])
            status = job.get("status")
            result = job.get("result") or {}

            if status != "done":
                record(f"{path.name}: ingest", FAIL,
                       f"status={status} result={result}")
                continue

            chunks = result.get("chunks", 0)
            pages = result.get("pages", 0)
            ocr_units = result.get("ocr_units", 0)
            duplicates = result.get("chunks_skipped_as_duplicate", 0)

            if chunks > 0:
                detail = f"{chunks} chunks from {pages} unit(s), ocr_units={ocr_units}"
                record(f"{path.name}: ingest", PASS, detail)
            elif duplicates:
                # Legitimate: identical content was already indexed. Handled
                # gracefully by design (run with --fresh to re-index).
                record(f"{path.name}: ingest", PASS,
                       f"all {duplicates} chunk(s) already indexed (dedup, not an error)")
            else:
                record(f"{path.name}: ingest", FAIL,
                       f"0 chunks and 0 duplicates from {pages} unit(s)")

            ingested.append((path.name, result))

        if not ingested:
            record("any document ingested", FAIL, "nothing to query against")
            return 1

        print()
        print("--- Store verification ---")
        check("all formats ingested", len(ingested) == len(files),
              f"{len(ingested)}/{len(files)}")

        # Items 7 + 8: the sparse index must cover every file, not just the
        # last one ingested, and must share ids with the dense store.
        sparse_total = None
        for _name, result in ingested:
            if result.get("sparse_index_total") is not None:
                sparse_total = result["sparse_index_total"]
        if sparse_total is None:
            record("sparse index populated", WARN, "worker did not report a sparse total")
        else:
            check("sparse index survived all ingests", sparse_total >= len(ingested),
                  f"{sparse_total} rows after ingesting {len(ingested)} files")

        try:
            import psycopg2

            conn = psycopg2.connect(
                host=os.getenv("POSTGRES_HOST", "localhost"),
                port=os.getenv("POSTGRES_PORT", "5432"),
                dbname=os.getenv("POSTGRES_DB", "rag_metadata"),
                user=os.getenv("POSTGRES_USER", "raglab"),
                password=os.getenv("POSTGRES_PASSWORD", "raglab"),
            )
            cur = conn.cursor()
            cur.execute("SELECT count(*), count(DISTINCT metadata->>'source') FROM chunks")
            rows, sources_count = cur.fetchone()
            check("sparse index has real sources (not placeholders)",
                  sources_count > 1,
                  f"{rows} rows across {sources_count} distinct source(s)")

            cur.execute("SELECT count(*) FROM chunks WHERE metadata->>'source' = 'test'")
            placeholder = cur.fetchone()[0]
            check("no hardcoded 'test' placeholder rows", placeholder == 0,
                  f"{placeholder} placeholder row(s)")

            cur.execute(
                "SELECT count(*) FROM chunks WHERE metadata->>'source' IS NOT NULL"
            )
            check("sparse rows carry source metadata", cur.fetchone()[0] == rows,
                  f"{rows} row(s) checked")

            cur.close()
            conn.close()
        except Exception as exc:
            record("sparse store inspection", WARN, f"{type(exc).__name__}: {exc}")

        # Dense vs sparse id overlap: the precondition for RRF to fuse rather
        # than double-count.
        try:
            scroll = requests.post(
                "http://localhost:6333/collections/rag_documents/points/scroll",
                json={"limit": 200, "with_payload": False, "with_vector": False},
                timeout=15,
            ).json()
            dense_ids = {str(p["id"]) for p in scroll["result"]["points"]}

            import psycopg2

            conn = psycopg2.connect(
                host=os.getenv("POSTGRES_HOST", "localhost"),
                port=os.getenv("POSTGRES_PORT", "5432"),
                dbname=os.getenv("POSTGRES_DB", "rag_metadata"),
                user=os.getenv("POSTGRES_USER", "raglab"),
                password=os.getenv("POSTGRES_PASSWORD", "raglab"),
            )
            cur = conn.cursor()
            cur.execute("SELECT id::text FROM chunks")
            sparse_ids = {row[0] for row in cur.fetchall()}
            cur.close()
            conn.close()

            overlap = dense_ids & sparse_ids
            check("dense and sparse stores share chunk ids", len(overlap) > 0,
                  f"{len(overlap)} shared id(s) of {len(dense_ids)} dense / "
                  f"{len(sparse_ids)} sparse")
        except Exception as exc:
            record("id overlap check", WARN, f"{type(exc).__name__}: {exc}")

        print()
        print("--- Ollama LLM availability ---")
        try:
            tags = requests.get("http://127.0.0.1:11434/api/tags", timeout=10).json()
            names = [m["name"] for m in tags.get("models", [])]
            check("llama3.1:8b present", any("llama3.1:8b" in n for n in names), str(names[:5]))
        except Exception as exc:
            record("Ollama reachable", WARN, f"{type(exc).__name__}: {exc}")

        print()
        print("--- Retrieval and answering ---")

        # An in-domain question with a specific factual answer.
        try:
            answer = query("What is the refund window for a standard item?")
            text = answer.get("answer", "")
            sources = answer.get("sources", [])
            contexts = answer.get("contexts", [])

            check("in-domain: answer returned", bool(text.strip()), repr(text[:200]))
            check("in-domain: grounded in retrieved context",
                  any("30 days" in c for c in contexts),
                  f"{len(contexts)} context(s) retrieved")
            check("in-domain: answer states 30 days", "30" in text,
                  repr(text[:200]))
            check("in-domain: sources returned", len(sources) > 0, f"{len(sources)} source(s)")

            # Item 4: citations must name the file (and page), not a bare UUID.
            import re as _re

            uuid_like = _re.compile(
                r"^\[?\d*\]?\s*[0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-"
                r"[0-9a-f]{4}-[0-9a-f]{12}$",
                _re.IGNORECASE,
            )
            bare_uuids = [s for s in sources if uuid_like.match(str(s).strip())]
            check("citations name a file, not a UUID", not bare_uuids,
                  f"sources = {sources[:3]}")

            refs = answer.get("source_refs") or []
            check("structured source_refs returned", len(refs) > 0,
                  f"{len(refs)} ref(s)")
            if refs:
                check("source_refs carry a filename",
                      all(r.get("source") and r["source"] != "unknown" for r in refs),
                      str([r.get("source") for r in refs][:4]))
                check("source_refs carry a location",
                      any(r.get("location") for r in refs),
                      str([r.get("location") for r in refs][:4]))
        except Exception as exc:
            record("in-domain query", FAIL, f"{type(exc).__name__}: {exc}")

        # Excel-only content: proves the second sheet reached the index.
        try:
            answer = query("What is the shipping lead time for Europe?")
            text = answer.get("answer", "")
            contexts = answer.get("contexts", [])
            check("excel sheet 2 retrieved",
                  any("5 business days" in c for c in contexts),
                  repr(text[:160]))
        except Exception as exc:
            record("excel sheet query", FAIL, f"{type(exc).__name__}: {exc}")

        # OCR-only content: proves scanned text reached the index.
        try:
            answer = query("Are refunds available within 30 days according to the scanned notice?")
            contexts = answer.get("contexts", [])
            check("OCR content retrieved",
                  any("REFUNDS WITHIN 30 DAYS" in c.upper() for c in contexts),
                  f"{len(contexts)} context(s)")
        except Exception as exc:
            record("OCR query", FAIL, f"{type(exc).__name__}: {exc}")

        # Out-of-scope question: the anti-hallucination requirement.
        try:
            answer = query("What is the capital of Mongolia?")
            text = answer.get("answer", "").lower()
            refused = any(
                marker in text
                for marker in ("don't know", "do not know", "not in the context",
                               "cannot find", "could not find", "no information")
            )
            if refused:
                record("out-of-scope refused", PASS, repr(text[:160]))
            else:
                record("out-of-scope refused", FAIL,
                       f"possible hallucination: {text[:200]!r}")
        except Exception as exc:
            record("out-of-scope query", FAIL, f"{type(exc).__name__}: {exc}")

        print()
        print("--- Other retrieval modes ---")
        for mode in ("HyDE", "Multi-Query", "Reranked"):
            try:
                answer = query("What is the refund window for a standard item?", mode=mode)
                check(f"mode {mode} answered", bool(answer.get("answer", "").strip()),
                      f"{len(answer.get('sources', []))} source(s)")
            except Exception as exc:
                record(f"mode {mode}", FAIL, f"{type(exc).__name__}: {str(exc)[:120]}")

    print()
    print("=" * 74)
    passed = sum(1 for _, s, _ in results if s == PASS)
    failed = sum(1 for _, s, _ in results if s == FAIL)
    warned = sum(1 for _, s, _ in results if s == WARN)
    print(f"RESULT: {passed} passed, {failed} failed, {warned} warnings")
    print("=" * 74)

    if failed:
        print("\nFailures:")
        for name, status, detail in results:
            if status == FAIL:
                print(f"  - {name}: {detail}")

    return 1 if failed else 0


if __name__ == "__main__":
    raise SystemExit(main())
