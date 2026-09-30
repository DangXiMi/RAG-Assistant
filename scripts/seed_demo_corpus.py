"""Build and ingest a small demo knowledge base.

Creates one document per supported format with mutually consistent content, so
the demo can show multi-format ingestion *and* answer questions that span the
files. The corpus is a fictional customer-support knowledge base, which matches
the questions the UI suggests.

Usage:
    python scripts/seed_demo_corpus.py            # create + ingest
    python scripts/seed_demo_corpus.py --no-ingest # create files only

Requires the API and worker to be running.
"""
from __future__ import annotations

import argparse
import sys
import time
from pathlib import Path

import requests

PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT))

DEMO_DIR = PROJECT_ROOT / "data" / "demo_documents"
BASE_URL = "http://127.0.0.1:8010"

# The refund window is stated consistently everywhere, so answers can be
# cross-checked between formats.
REFUND_WINDOW = "30 days from delivery"

REFUND_POLICY_TEXT = f"""Refund Policy

We want you to be completely satisfied with your purchase.

Standard items may be returned for a full refund within {REFUND_WINDOW}.
The item must be unused and in its original packaging.

Sale items are not refundable, except where the item arrived damaged or faulty.

Shipping costs are non-refundable for change-of-mind returns. If the return is
due to our error, we refund the original shipping cost as well.

To start a return, contact support with your order number. Refunds are issued
to the original payment method within 5 business days of us receiving the item.
"""

SHIPPING_TEXT = """Shipping Policy

Orders placed before 2pm local time are dispatched the same working day.

Delivery lead times:
- Europe: 5 business days
- North America: 7 business days
- Rest of world: 12 business days

Express shipping is available at checkout for an additional fee. Express
shipping costs are non-refundable.

We do not ship to PO boxes for orders over 500 EUR.
"""

WARRANTY_TEXT = """Warranty Terms

All electronics carry a 24-month manufacturer warranty covering defects in
materials and workmanship.

The warranty does not cover accidental damage, liquid damage, or normal wear.

Warranty claims are handled separately from refunds. A warranty replacement
does not extend the original warranty period.
"""


def build_files(directory: Path) -> list[Path]:
    directory.mkdir(parents=True, exist_ok=True)
    files: list[Path] = []

    # 1. Plain text.
    path = directory / "refund_policy.txt"
    path.write_text(REFUND_POLICY_TEXT, encoding="utf-8")
    files.append(path)

    # 2. Markdown.
    path = directory / "warranty_terms.md"
    path.write_text("# " + WARRANTY_TEXT, encoding="utf-8")
    files.append(path)

    # 3. Word document with a real table.
    from docx import Document

    path = directory / "refund_matrix.docx"
    document = Document()
    document.add_heading("Refund Matrix", level=1)
    document.add_paragraph(
        "The table below summarises which items can be refunded and the window."
    )
    table = document.add_table(rows=4, cols=2)
    rows = [
        ("Item type", "Refund window"),
        ("Standard", REFUND_WINDOW),
        ("Sale item", "Not refundable"),
        ("Damaged or faulty", "Full refund, any time"),
    ]
    for i, (left, right) in enumerate(rows):
        table.cell(i, 0).text = left
        table.cell(i, 1).text = right
    document.add_paragraph("Contact support to begin a return.")
    document.save(path)
    files.append(path)

    # 4. Excel with two sheets, so a question answered only by sheet 2 works.
    from openpyxl import Workbook

    path = directory / "service_levels.xlsx"
    workbook = Workbook()
    workbook.remove(workbook.active)

    refunds = workbook.create_sheet("Refunds")
    refunds.append(["Item type", "Refund window"])
    refunds.append(["Standard", REFUND_WINDOW])
    refunds.append(["Sale item", "Not refundable"])

    shipping = workbook.create_sheet("Shipping")
    shipping.append(["Region", "Lead time"])
    shipping.append(["Europe", "5 business days"])
    shipping.append(["North America", "7 business days"])
    shipping.append(["Rest of world", "12 business days"])

    workbook.save(path)
    files.append(path)

    # 5. CSV.
    path = directory / "returns_faq.csv"
    path.write_text(
        "Question,Answer\n"
        f"How long do I have to return an item?,{REFUND_WINDOW}\n"
        "Are sale items refundable?,No\n"
        "How long do refunds take?,5 business days after we receive the item\n",
        encoding="utf-8",
    )
    files.append(path)

    # 6. HTML.
    path = directory / "support_hours.html"
    path.write_text(
        "<html><body>"
        "<h1>Contacting Support</h1>"
        "<p>Our support team is available Monday to Friday, 9am to 6pm CET.</p>"
        "<table><tr><th>Channel</th><th>Response time</th></tr>"
        "<tr><td>Email</td><td>24 hours</td></tr>"
        "<tr><td>Live chat</td><td>Under 5 minutes</td></tr></table>"
        "</body></html>",
        encoding="utf-8",
    )
    files.append(path)

    # 7. PDF with a text layer and two pages, so page citations are visible.
    import pymupdf

    path = directory / "terms_and_conditions.pdf"
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 90), "Terms and Conditions", fontsize=16)
    page.insert_text(
        (72, 130),
        f"Standard items may be returned within {REFUND_WINDOW}.",
        fontsize=12,
    )
    page.insert_text(
        (72, 155), "Sale items are not refundable.", fontsize=12
    )
    second = document.new_page()
    second.insert_text((72, 90), "Page two: Shipping", fontsize=16)
    second.insert_text(
        (72, 130), "Europe delivery takes 5 business days.", fontsize=12
    )
    second.insert_text(
        (72, 155), "Express shipping is non-refundable.", fontsize=12
    )
    document.save(str(path))
    document.close()
    files.append(path)

    # 8. A scanned-looking image, to demonstrate OCR end to end.
    from PIL import Image, ImageDraw

    path = directory / "scanned_notice.png"
    image = Image.new("RGB", (1200, 320), "white")
    draw = ImageDraw.Draw(image)
    draw.text((30, 60), "IMPORTANT NOTICE", fill="black", font_size=54)
    draw.text((30, 150), "REFUNDS WITHIN 30 DAYS OF DELIVERY", fill="black", font_size=42)
    draw.text((30, 230), "KEEP YOUR RECEIPT", fill="black", font_size=42)
    image.save(path)
    files.append(path)

    return files


def ingest(path: Path, base_url: str = BASE_URL) -> dict:
    with open(path, "rb") as handle:
        response = requests.post(
            f"{base_url}/api/v1/ingest",
            files={"file": (path.name, handle, "application/octet-stream")},
            timeout=120,
        )
    response.raise_for_status()
    return response.json()


def wait_for_job(job_id: str, base_url: str = BASE_URL, timeout: int = 300) -> dict:
    deadline = time.time() + timeout
    while time.time() < deadline:
        response = requests.get(f"{base_url}/api/v1/job/{job_id}", timeout=30)
        if response.status_code == 200:
            data = response.json()
            if data.get("status") in ("done", "failed"):
                return data
        time.sleep(2)
    return {"status": "timeout", "result": None}


def reset_stores() -> None:
    """Empty both retrieval stores before seeding.

    Without this, leftovers from tests and earlier evaluations stay in the
    index and get retrieved as context, so the assistant cites documents that
    are not part of the demo (observed as "e2e_test" and "jwst_sample.txt"
    appearing in the source list).
    """
    from qdrant_client import QdrantClient

    from src.config.config import CONFIG
    from src.ingestion import sparse_store

    client = QdrantClient(host="localhost", port=6333)
    name = CONFIG["qdrant"]["collection_name"]
    if client.collection_exists(name):
        client.delete_collection(name)
        print(f"Cleared dense collection '{name}'")
    client.close()

    conn = sparse_store.connect()
    try:
        # Truncate rather than drop: the schema is unchanged, and rebuilding it
        # is not the point of a data reset.
        with conn.cursor() as cur:
            cur.execute("DELETE FROM chunks")
        print("Cleared sparse index")
    finally:
        conn.close()


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--no-ingest", action="store_true", help="only create files")
    parser.add_argument(
        "--keep-existing",
        action="store_true",
        help="do not clear the stores first (appends instead)",
    )
    parser.add_argument("--base-url", default=BASE_URL)
    args = parser.parse_args()

    base_url = args.base_url

    files = build_files(DEMO_DIR)
    print(f"Created {len(files)} demo document(s) in {DEMO_DIR}:")
    for path in files:
        print(f"  {path.name:<30} {path.stat().st_size:>8} bytes")

    if args.no_ingest:
        return 0

    if not args.keep_existing:
        print()
        reset_stores()

    # Clear the dedup cache so a re-run re-indexes rather than reporting
    # everything as already present.
    try:
        import redis as redis_lib

        client = redis_lib.Redis(host="localhost", port=6379, decode_responses=True)
        keys = list(client.scan_iter("chunk:hash:*", count=1000))
        if keys:
            client.delete(*keys)
            print(f"Cleared {len(keys)} dedup key(s)")
    except Exception as exc:
        print(f"Could not clear dedup cache ({type(exc).__name__}); continuing")

    print(f"\nIngesting through {base_url} ...")
    failures = 0
    for path in files:
        try:
            job = wait_for_job(ingest(path, base_url)["job_id"], base_url)
        except Exception as exc:
            print(f"  [!] {path.name}: {type(exc).__name__}: {exc}")
            failures += 1
            continue

        result = job.get("result") or {}
        if job.get("status") == "done":
            chunks = result.get("chunks", 0)
            if chunks:
                print(f"  [ok] {path.name:<30} {chunks} chunk(s)")
            else:
                print(
                    f"  [ok] {path.name:<30} already indexed "
                    f"({result.get('chunks_skipped_as_duplicate', 0)} duplicate)"
                )
        else:
            failures += 1
            print(f"  [XX] {path.name:<30} {result.get('error')}")

    print()
    if failures:
        print(f"{failures} document(s) failed to ingest")
        return 1

    print("Demo corpus ready. Ask: \"What is the refund policy?\"")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
