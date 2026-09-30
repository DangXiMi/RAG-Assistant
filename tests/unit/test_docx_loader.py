# tests/unit/test_docx_loader.py
"""Tests for Word .docx loading, especially table preservation.

The previous implementation read only `doc.paragraphs` and silently discarded
every table; these tests lock in that tables now reach the index.
"""
import pytest

pytest.importorskip("docx")

from docx import Document  # noqa: E402

from src.ingestion.loaders import load_document  # noqa: E402


@pytest.fixture
def simple_docx(tmp_path):
    path = tmp_path / "policy.docx"
    document = Document()
    document.add_heading("Refund Policy", level=1)
    document.add_paragraph("Refunds are processed within 30 days of delivery.")
    document.add_paragraph("Sale items are not refundable.")
    document.save(path)
    return path


@pytest.fixture
def docx_with_table(tmp_path):
    path = tmp_path / "refunds.docx"
    document = Document()
    document.add_heading("Refund Matrix", level=1)
    document.add_paragraph("The table below summarises the refund windows.")

    table = document.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "Item type"
    table.cell(0, 1).text = "Refund window"
    table.cell(1, 0).text = "Standard"
    table.cell(1, 1).text = "30 days"
    table.cell(2, 0).text = "Sale item"
    table.cell(2, 1).text = "Not refundable"

    document.add_paragraph("Contact support for exceptions.")
    document.save(path)
    return path


def test_loads_paragraph_text(simple_docx):
    pages = load_document(simple_docx)

    assert len(pages) == 1
    assert "30 days of delivery" in pages[0].text
    assert pages[0].metadata["source"] == "policy.docx"


def test_headings_become_markdown(simple_docx):
    pages = load_document(simple_docx)

    assert "# Refund Policy" in pages[0].text


def test_table_content_is_preserved(docx_with_table):
    pages = load_document(docx_with_table)

    # This is the regression that mattered: table cells used to vanish.
    assert "30 days" in pages[0].text
    assert "Not refundable" in pages[0].text
    assert "Refund window" in pages[0].text


def test_table_is_rendered_as_markdown(docx_with_table):
    pages = load_document(docx_with_table)

    assert "| Item type" in pages[0].text
    assert "|---" in pages[0].text


def test_table_count_is_recorded(docx_with_table):
    pages = load_document(docx_with_table)

    assert pages[0].metadata["tables"] == 1


def test_tables_appear_in_document_order(docx_with_table):
    pages = load_document(docx_with_table)
    text = pages[0].text

    intro = text.index("summarises the refund windows")
    table = text.index("| Item type")
    outro = text.index("Contact support for exceptions")

    assert intro < table < outro


def test_no_tables_records_zero(simple_docx):
    pages = load_document(simple_docx)

    assert pages[0].metadata["tables"] == 0


def test_headers_and_footers_are_captured(tmp_path):
    path = tmp_path / "headed.docx"
    document = Document()
    document.add_paragraph("Main body text.")
    document.sections[0].header.paragraphs[0].text = "Internal Use Only"
    document.save(path)

    pages = load_document(path)

    assert "Main body text." in pages[0].text
    assert "Internal Use Only" in pages[0].text


def test_chunking_a_docx_keeps_source_metadata(docx_with_table):
    from src.ingestion.chunker import chunk_pages

    pages = load_document(docx_with_table)
    chunks = chunk_pages(pages)

    assert chunks
    for chunk in chunks:
        assert chunk.metadata["source"] == "refunds.docx"
