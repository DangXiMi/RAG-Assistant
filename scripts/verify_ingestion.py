"""End-to-end check of the multi-format ingestion pipeline.

Generates a synthetic document in every supported format, runs it through
`load_document` -> `chunk_pages`, and asserts that the expected content
reached the chunks along with usable citation metadata.

Usage:
    python scripts/verify_ingestion.py
"""
from __future__ import annotations

import sys
import tempfile
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT))

from src.ingestion.chunker import chunk_pages  # noqa: E402
from src.ingestion.loaders import load_document  # noqa: E402
from src.ingestion.loaders.capabilities import (  # noqa: E402
    log_summary,
    ocr_available,
    ocr_unavailable_reason,
)

PASS = "PASS"
FAIL = "FAIL"
SKIP = "SKIP"

results: list[tuple[str, str, str]] = []


def record(name: str, status: str, detail: str = "") -> None:
    results.append((name, status, detail))
    icon = {PASS: "[ok]", FAIL: "[XX]", SKIP: "[--]"}[status]
    print(f"{icon} {name}" + (f" -- {detail}" if detail else ""))


def check(name: str, condition: bool, detail: str = "") -> None:
    record(name, PASS if condition else FAIL, detail)


def chunks_text(chunks) -> str:
    return " ".join(c.text for c in chunks)


# --- Fixtures ---------------------------------------------------------------

def make_text(path: Path) -> Path:
    path.write_text(
        "Refund Policy\n\nRefunds are processed within 30 days of delivery.\n",
        encoding="utf-8",
    )
    return path


def make_html(path: Path) -> Path:
    path.write_text(
        "<html><body><h1>Refund Policy</h1>"
        "<p>Refunds are processed within 30 days.</p>"
        "<table><tr><th>Item</th><th>Window</th></tr>"
        "<tr><td>Sale item</td><td>Not refundable</td></tr></table>"
        "</body></html>",
        encoding="utf-8",
    )
    return path


def make_csv(path: Path) -> Path:
    path.write_text(
        "Item type,Refund window\nStandard,30 days\nSale item,Not refundable\n",
        encoding="utf-8",
    )
    return path


def make_docx(path: Path) -> Path:
    from docx import Document

    document = Document()
    document.add_heading("Refund Policy", level=1)
    document.add_paragraph("Refunds are processed within 30 days of delivery.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Item type"
    table.cell(0, 1).text = "Refund window"
    table.cell(1, 0).text = "Sale item"
    table.cell(1, 1).text = "Not refundable"
    document.save(path)
    return path


def make_xlsx(path: Path) -> Path:
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.remove(workbook.active)
    policy = workbook.create_sheet("Policy")
    policy.append(["Item type", "Refund window"])
    policy.append(["Standard", "30 days"])
    policy.append(["Sale item", "Not refundable"])
    shipping = workbook.create_sheet("Shipping")
    shipping.append(["Region", "Lead time"])
    shipping.append(["Europe", "5 business days"])
    workbook.save(path)
    return path


def make_pdf(path: Path) -> Path:
    import pymupdf

    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 100), "Refunds are processed within 30 days.", fontsize=12)
    second = document.new_page()
    second.insert_text((72, 100), "Sale items are not refundable.", fontsize=12)
    document.save(str(path))
    document.close()
    return path


def make_image(path: Path) -> Path:
    from PIL import Image, ImageDraw

    image = Image.new("RGB", (1000, 220), "white")
    ImageDraw.Draw(image).text(
        (20, 90), "REFUNDS WITHIN 30 DAYS", fill="black", font_size=52
    )
    image.save(path)
    return path


def make_scanned_pdf(path: Path) -> Path:
    import pymupdf

    image_path = path.with_suffix(".png")
    make_image(image_path)

    document = pymupdf.open()
    page = document.new_page(width=1000, height=220)
    page.insert_image(pymupdf.Rect(0, 0, 1000, 220), filename=str(image_path))
    document.save(str(path))
    document.close()
    return path


# --- Checks -----------------------------------------------------------------

def check_format(workdir: Path, filename: str, builder, expected: list[str],
                 expect_metadata: dict | None = None, needs_ocr: bool = False) -> None:
    path = builder(workdir / filename)

    if needs_ocr and not ocr_available():
        record(f"{filename}: OCR path", SKIP, ocr_unavailable_reason())
        return

    try:
        pages = load_document(path)
    except Exception as exc:
        record(f"{filename}: load", FAIL, f"{type(exc).__name__}: {exc}")
        return

    record(f"{filename}: load", PASS, f"{len(pages)} unit(s)")

    if not pages:
        record(f"{filename}: content", FAIL, "no units extracted")
        return

    chunks = chunk_pages(pages)
    text = chunks_text(chunks)

    for needle in expected:
        check(f"{filename}: contains {needle!r}", needle.lower() in text.lower())

    if expect_metadata:
        for key, value in expect_metadata.items():
            found = {c.metadata.get(key) for c in chunks}
            check(
                f"{filename}: metadata {key}={value!r}",
                value in found,
                f"found {sorted(str(v) for v in found)}",
            )

    # Every chunk must be attributable to a source file for citations.
    sources = {c.metadata.get("source") for c in chunks}
    check(
        f"{filename}: citation source present",
        sources == {filename},
        f"{sources}",
    )


def main() -> int:
    print("=" * 72)
    print("Ingestion capability summary")
    print("=" * 72)
    log_summary()
    for name, available in __import__(
        "src.ingestion.loaders.capabilities", fromlist=["summary"]
    ).summary().items():
        print(f"  {name:<26} {'available' if available else 'MISSING'}")

    print()
    print("=" * 72)
    print("Format checks")
    print("=" * 72)

    with tempfile.TemporaryDirectory() as tmp:
        workdir = Path(tmp)

        check_format(
            workdir, "policy.txt", make_text,
            expected=["30 days"],
        )
        check_format(
            workdir, "policy.html", make_html,
            expected=["30 days", "Not refundable"],
        )
        check_format(
            workdir, "refunds.csv", make_csv,
            expected=["30 days", "Not refundable"],
        )
        check_format(
            workdir, "policy.docx", make_docx,
            expected=["30 days", "Not refundable"],
        )
        check_format(
            workdir, "refunds.xlsx", make_xlsx,
            expected=["30 days", "Not refundable", "5 business days"],
            expect_metadata={"sheet": "Shipping"},
        )
        check_format(
            workdir, "policy.pdf", make_pdf,
            expected=["30 days", "not refundable"],
            expect_metadata={"page": 2},
        )
        check_format(
            workdir, "scan.png", make_image,
            expected=["REFUND", "30"],
            needs_ocr=True,
        )
        check_format(
            workdir, "scanned.pdf", make_scanned_pdf,
            expected=["REFUND", "30"],
            needs_ocr=True,
        )

    print()
    print("=" * 72)
    passed = sum(1 for _, s, _ in results if s == PASS)
    failed = sum(1 for _, s, _ in results if s == FAIL)
    skipped = sum(1 for _, s, _ in results if s == SKIP)
    print(f"RESULT: {passed} passed, {failed} failed, {skipped} skipped")
    print("=" * 72)

    if failed:
        print("\nFailures:")
        for name, status, detail in results:
            if status == FAIL:
                print(f"  - {name}: {detail}")

    return 1 if failed else 0


if __name__ == "__main__":
    raise SystemExit(main())
